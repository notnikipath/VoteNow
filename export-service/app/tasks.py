from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from celery import Celery
from fpdf import FPDF
from openpyxl import Workbook

from shared.votenow_common.settings import Settings

settings = Settings(service_name="export-service", service_port=8505)
celery_app = Celery("votenow-export", broker=settings.redis_url, backend=settings.redis_url)
celery_app.conf.update(
    task_always_eager=settings.celery_task_always_eager,
    task_serializer="json",
    result_serializer="json",
)

export_root = Path(settings.export_root)
export_root.mkdir(parents=True, exist_ok=True)


def _timestamp() -> str:
    return datetime.now(UTC).strftime("%Y%m%d%H%M%S")


@celery_app.task(name="votenow.generate_excel")
def generate_excel_report(survey_id: str, table_payload: dict[str, Any]) -> dict[str, str]:
    filename = f"{survey_id}-{_timestamp()}.xlsx"
    path = export_root / filename

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "VoteNow"
    sheet.append(["Survey ID", survey_id])
    sheet.append(["Generated At", datetime.now(UTC).isoformat()])
    sheet.append(["Total Responses", table_payload.get("total_responses", 0)])
    sheet.append([])
    sheet.append(["Question ID", "Title", "Type", "Data"])

    for row in table_payload.get("rows", []):
        sheet.append(
            [
                row.get("question_id"),
                row.get("title"),
                row.get("type"),
                str(row.get("payload", {})),
            ]
        )

    workbook.save(path)
    return {"status": "completed", "format": "excel", "file_name": filename, "file_path": str(path)}


@celery_app.task(name="votenow.generate_pdf")
def generate_pdf_report(survey_id: str, table_payload: dict[str, Any]) -> dict[str, str]:
    filename = f"{survey_id}-{_timestamp()}.pdf"
    path = export_root / filename

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, f"VoteNow Report: {survey_id}", ln=1)
    pdf.set_font("Helvetica", size=11)
    pdf.cell(0, 8, f"Generated: {datetime.now(UTC).isoformat()}", ln=1)
    pdf.cell(0, 8, f"Total responses: {table_payload.get('total_responses', 0)}", ln=1)
    pdf.ln(3)

    for row in table_payload.get("rows", []):
        pdf.set_font("Helvetica", "B", 11)
        pdf.multi_cell(0, 7, f"{row.get('question_id')} - {row.get('title')}")
        pdf.set_font("Helvetica", size=10)
        pdf.multi_cell(0, 6, f"Type: {row.get('type')}")
        pdf.multi_cell(0, 6, f"Data: {row.get('payload')}")
        pdf.ln(2)

    pdf.output(str(path))
    return {"status": "completed", "format": "pdf", "file_name": filename, "file_path": str(path)}
